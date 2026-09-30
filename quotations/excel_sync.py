import io
import os
import openpyxl
from django.conf import settings
from quotations.models import Customer, Product
from quotations.management.commands.import_customers import load_workbook_safely, read_file_safely

def get_excel_file_path(filename='customer address raj.xlsx'):
    return os.path.join(settings.BASE_DIR, filename)

def append_or_update_customer_in_excel(customer, original_name=None, filename='customer address raj.xlsx'):
    """
    Appends a new customer or updates an existing customer directly in the Excel file.
    If original_name is provided, it first matches by the previous name in case name was edited.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        # Find header row
        header_row_idx = None
        col_indices = {}

        for r in range(1, min(15, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'PARTY NAME' in val:
                    header_row_idx = r
                    break
            if header_row_idx:
                for c_idx, val in enumerate(row_vals):
                    if 'PARTY NAME' in val:
                        col_indices['name'] = c_idx + 1
                    elif 'EMAIL' in val:
                        col_indices['email'] = c_idx + 1
                    elif 'MOBILE' in val or 'PHONE' in val:
                        col_indices['mobile'] = c_idx + 1
                    elif 'GST' in val:
                        col_indices['gstin'] = c_idx + 1
                    elif 'ADDRESS' in val:
                        col_indices['address'] = c_idx + 1
                break

        if not header_row_idx or 'name' not in col_indices:
            # Fallback to standard columns if not found
            header_row_idx = 2
            col_indices = {'name': 1, 'email': 2, 'mobile': 3, 'gstin': 4, 'address': 5}

        # Check if customer already exists in sheet
        target_row = None
        orig_name_lower = original_name.strip().lower() if original_name else None
        target_name_lower = customer.name.strip().lower()
        target_gstin_clean = customer.gstin.strip().upper() if customer.gstin and customer.gstin.strip().upper() not in ('', 'NIL', 'N/A') else None

        # 1. Search by original_name if provided
        if orig_name_lower:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['name']).value
                if val and str(val).strip().lower() == orig_name_lower:
                    target_row = r
                    break

        # 2. Search by current customer name
        if not target_row:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['name']).value
                if val and str(val).strip().lower() == target_name_lower:
                    target_row = r
                    break

        # 3. Fallback: match by GSTIN if available
        if not target_row and target_gstin_clean and 'gstin' in col_indices:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['gstin']).value
                if val and str(val).strip().upper() == target_gstin_clean:
                    target_row = r
                    break

        is_new_row = False
        if not target_row:
            target_row = sheet.max_row + 1
            is_new_row = True

        # Write values
        sheet.cell(row=target_row, column=col_indices['name'], value=customer.name.strip())
        sheet.cell(row=target_row, column=col_indices['email'], value=customer.email.strip() if customer.email else 'Nil')
        sheet.cell(row=target_row, column=col_indices['mobile'], value=customer.mobile.strip() if customer.mobile else 'Nil')
        sheet.cell(row=target_row, column=col_indices['gstin'], value=customer.gstin.strip() if customer.gstin else 'Nil')
        sheet.cell(row=target_row, column=col_indices['address'], value=customer.address.strip() if customer.address else '')

        # Attempt to save to disk
        wb.save(file_path)
        action_desc = "added to" if is_new_row else "updated in"
        return True, f"Successfully {action_desc} '{filename}' (Row {target_row})!"

    except PermissionError:
        return False, (
            f"Saved to database, but '{filename}' is currently open in Microsoft Excel. "
            f"Please save & close Excel, then click 'Sync to Excel' to write this record."
        )
    except Exception as e:
        return False, f"Error saving to Excel: {str(e)}"


def sync_all_customers_to_excel(filename='customer address raj.xlsx'):
    """
    Synchronizes all customers from the database into the Excel file.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        header_row_idx = 5
        col_indices = {'name': 1, 'email': 2, 'mobile': 3, 'gstin': 4, 'address': 5}

        for r in range(1, min(15, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'PARTY NAME' in val:
                    header_row_idx = r
                    for c_idx2, val2 in enumerate(row_vals):
                        if 'PARTY NAME' in val2:
                            col_indices['name'] = c_idx2 + 1
                        elif 'EMAIL' in val2:
                            col_indices['email'] = c_idx2 + 1
                        elif 'MOBILE' in val2 or 'PHONE' in val2:
                            col_indices['mobile'] = c_idx2 + 1
                        elif 'GST' in val2:
                            col_indices['gstin'] = c_idx2 + 1
                        elif 'ADDRESS' in val2:
                            col_indices['address'] = c_idx2 + 1
                    break
            if header_row_idx != 5:
                break

        # Map existing rows by customer name lowercase
        existing_rows = {}
        for r in range(header_row_idx + 1, sheet.max_row + 1):
            val = sheet.cell(row=r, column=col_indices['name']).value
            if val and str(val).strip():
                existing_rows[str(val).strip().lower()] = r

        customers = Customer.objects.filter(is_active=True).order_by('id')
        current_max = sheet.max_row

        for cust in customers:
            key = cust.name.strip().lower()
            if key in existing_rows:
                r_idx = existing_rows[key]
            else:
                current_max += 1
                r_idx = current_max
                existing_rows[key] = r_idx

            sheet.cell(row=r_idx, column=col_indices['name'], value=cust.name.strip())
            sheet.cell(row=r_idx, column=col_indices['email'], value=cust.email.strip() if cust.email else 'Nil')
            sheet.cell(row=r_idx, column=col_indices['mobile'], value=cust.mobile.strip() if cust.mobile else 'Nil')
            sheet.cell(row=r_idx, column=col_indices['gstin'], value=cust.gstin.strip() if cust.gstin else 'Nil')
            sheet.cell(row=r_idx, column=col_indices['address'], value=cust.address.strip() if cust.address else '')

        wb.save(file_path)
        return True, f"Successfully synchronized {customers.count()} customers to '{filename}'!"

    except PermissionError:
        return False, (
            f"Microsoft Excel currently has '{filename}' open. "
            f"Please close Excel and click 'Sync to Excel' again."
        )
    except Exception as e:
        return False, f"Failed to sync to Excel: {str(e)}"


def delete_customer_from_excel(customer_name, gstin=None, filename='customer address raj.xlsx'):
    """
    Deletes a customer row directly from customer address raj.xlsx.
    Matches by party name (case-insensitive) or GSTIN.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        header_row_idx = None
        col_indices = {}

        for r in range(1, min(15, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'PARTY NAME' in val:
                    header_row_idx = r
                    break
            if header_row_idx:
                for c_idx, val in enumerate(row_vals):
                    if 'PARTY NAME' in val:
                        col_indices['name'] = c_idx + 1
                    elif 'EMAIL' in val:
                        col_indices['email'] = c_idx + 1
                    elif 'MOBILE' in val or 'PHONE' in val:
                        col_indices['mobile'] = c_idx + 1
                    elif 'GST' in val:
                        col_indices['gstin'] = c_idx + 1
                    elif 'ADDRESS' in val:
                        col_indices['address'] = c_idx + 1
                break

        if not header_row_idx or 'name' not in col_indices:
            header_row_idx = 2
            col_indices = {'name': 1, 'email': 2, 'mobile': 3, 'gstin': 4, 'address': 5}

        target_row = None
        name_lower = customer_name.strip().lower()
        gstin_clean = gstin.strip().upper() if gstin and gstin.strip().upper() not in ('', 'NIL', 'N/A') else None

        # 1. Search by customer name
        for r in range(header_row_idx + 1, sheet.max_row + 1):
            val = sheet.cell(row=r, column=col_indices['name']).value
            if val and str(val).strip().lower() == name_lower:
                target_row = r
                break

        # 2. Fallback search by GSTIN
        if not target_row and gstin_clean and 'gstin' in col_indices:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['gstin']).value
                if val and str(val).strip().upper() == gstin_clean:
                    target_row = r
                    break

        if target_row:
            sheet.delete_rows(target_row, 1)
            wb.save(file_path)
            return True, f"Successfully removed '{customer_name}' from '{filename}'!"
        else:
            return True, f"Customer '{customer_name}' removed from database."

    except PermissionError:
        return False, (
            f"Removed from database, but '{filename}' is currently open in Microsoft Excel. "
            f"Please close Excel to finalize deletion from file."
        )
    except Exception as e:
        return False, f"Error deleting customer from Excel: {str(e)}"


def get_product_excel_file_path(filename='Product List.xlsx'):
    return os.path.join(settings.BASE_DIR, filename)


def append_or_update_product_in_excel(product, original_code=None, filename='Product List.xlsx'):
    """
    Appends a new product or updates an existing product directly in Product List.xlsx.
    If original_code is provided, matches by the previous code in case model_code was edited.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_product_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        header_row_idx = None
        col_indices = {'model_code': 1, 'category': 2, 'description': 3}

        for r in range(1, min(10, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'STOCK ITEM' in val or 'ITEM NAME' in val or 'MODEL' in val or 'PRODUCT' in val:
                    header_row_idx = r
                    break
            if header_row_idx:
                for c_idx, val in enumerate(row_vals):
                    if 'STOCK ITEM' in val or 'ITEM NAME' in val or 'MODEL' in val or 'PRODUCT' in val:
                        col_indices['model_code'] = c_idx + 1
                    elif 'BRAND' in val or 'MAKE' in val or 'CATEGORY' in val:
                        col_indices['category'] = c_idx + 1
                    elif 'DESC' in val:
                        col_indices['description'] = c_idx + 1
                    elif 'RATE' in val or 'PRICE' in val:
                        col_indices['unit_rate'] = c_idx + 1
                    elif 'UNIT' in val or 'UOM' in val:
                        col_indices['unit'] = c_idx + 1
                    elif 'HSN' in val:
                        col_indices['hsn_code'] = c_idx + 1
                    elif 'GST' in val:
                        col_indices['gst_rate'] = c_idx + 1
                break

        if not header_row_idx:
            header_row_idx = 2

        target_row = None
        orig_code_lower = original_code.strip().lower() if original_code else None
        target_code_lower = product.model_code.strip().lower()

        # 1. Search by original_code if provided
        if orig_code_lower:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['model_code']).value
                if val and str(val).strip().lower() == orig_code_lower:
                    target_row = r
                    break

        # 2. Search by current model_code
        if not target_row:
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                val = sheet.cell(row=r, column=col_indices['model_code']).value
                if val and str(val).strip().lower() == target_code_lower:
                    target_row = r
                    break

        is_new_row = False
        if not target_row:
            last_used_row = header_row_idx
            for r in range(header_row_idx + 1, sheet.max_row + 1):
                if any(sheet.cell(row=r, column=c).value is not None and str(sheet.cell(row=r, column=c).value).strip() != '' for c in range(1, max(4, sheet.max_column + 1))):
                    last_used_row = r
            target_row = last_used_row + 1
            is_new_row = True

        sheet.cell(row=target_row, column=col_indices['model_code'], value=product.model_code.strip())
        sheet.cell(row=target_row, column=col_indices['category'], value=product.category.strip() if product.category else 'GENERAL')
        sheet.cell(row=target_row, column=col_indices['description'], value=product.description.strip() if product.description else '')

        if 'unit' in col_indices and product.unit:
            sheet.cell(row=target_row, column=col_indices['unit'], value=str(product.unit).strip())
        if 'unit_rate' in col_indices and product.unit_rate is not None:
            sheet.cell(row=target_row, column=col_indices['unit_rate'], value=float(product.unit_rate))
        if 'hsn_code' in col_indices and product.hsn_code:
            sheet.cell(row=target_row, column=col_indices['hsn_code'], value=str(product.hsn_code).strip())
        if 'gst_rate' in col_indices and product.gst_rate is not None:
            sheet.cell(row=target_row, column=col_indices['gst_rate'], value=float(product.gst_rate))

        wb.save(file_path)
        action_desc = "added to" if is_new_row else "updated in"
        return True, f"Successfully {action_desc} '{filename}' (Row {target_row})!"

    except PermissionError:
        return False, (
            f"Saved to database, but '{filename}' is currently open in Microsoft Excel. "
            f"Please save & close Excel to update the spreadsheet."
        )
    except Exception as e:
        return False, f"Error saving product to Excel: {str(e)}"


def delete_product_from_excel(model_code, category=None, filename='Product List.xlsx'):
    """
    Deletes a product row directly from Product List.xlsx.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_product_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        header_row_idx = None
        col_indices = {'model_code': 1, 'category': 2}

        for r in range(1, min(10, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'STOCK ITEM' in val or 'ITEM NAME' in val or 'MODEL' in val or 'PRODUCT' in val:
                    header_row_idx = r
                    col_indices['model_code'] = c_idx + 1
                    break
            if header_row_idx:
                for c_idx, val in enumerate(row_vals):
                    if 'BRAND' in val or 'MAKE' in val or 'CATEGORY' in val:
                        col_indices['category'] = c_idx + 1
                break

        if not header_row_idx:
            header_row_idx = 2

        target_row = None
        code_lower = model_code.strip().lower()

        for r in range(header_row_idx + 1, sheet.max_row + 1):
            val = sheet.cell(row=r, column=col_indices['model_code']).value
            if val and str(val).strip().lower() == code_lower:
                if category:
                    cat_val = sheet.cell(row=r, column=col_indices['category']).value
                    if cat_val and str(cat_val).strip().lower() != category.strip().lower():
                        continue
                target_row = r
                break

        if target_row:
            sheet.delete_rows(target_row, 1)
            wb.save(file_path)
            return True, f"Successfully removed '{model_code}' from '{filename}'!"
        else:
            return True, f"Product '{model_code}' removed from database."

    except PermissionError:
        return False, (
            f"Removed from database, but '{filename}' is currently open in Microsoft Excel. "
            f"Please close Excel to finalize deletion from file."
        )
    except Exception as e:
        return False, f"Error deleting product from Excel: {str(e)}"


def sync_all_products_to_excel(filename='Product List.xlsx'):
    """
    Synchronizes all active products from the database into Product List.xlsx.
    Returns tuple: (success: bool, message: str)
    """
    file_path = get_product_excel_file_path(filename)
    if not os.path.exists(file_path):
        return False, f"Excel file not found at: {file_path}"

    try:
        wb = load_workbook_safely(file_path)
        sheet = wb.active

        header_row_idx = 2
        col_indices = {'model_code': 1, 'category': 2, 'description': 3}

        for r in range(1, min(10, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(row=r, column=c).value or '').strip().upper() for c in range(1, sheet.max_column + 1)]
            for c_idx, val in enumerate(row_vals):
                if 'STOCK ITEM' in val or 'ITEM NAME' in val:
                    header_row_idx = r
                    col_indices['model_code'] = c_idx + 1
                    for c_idx2, val2 in enumerate(row_vals):
                        if 'BRAND' in val2 or 'MAKE' in val2:
                            col_indices['category'] = c_idx2 + 1
                        elif 'DESC' in val2:
                            col_indices['description'] = c_idx2 + 1
                    break
            if header_row_idx != 2:
                break

        # Map existing rows by model_code lowercase
        existing_rows = {}
        for r in range(header_row_idx + 1, sheet.max_row + 1):
            val = sheet.cell(row=r, column=col_indices['model_code']).value
            if val and str(val).strip():
                existing_rows[str(val).strip().lower()] = r

        products = Product.objects.filter(is_active=True).order_by('id')
        current_max = sheet.max_row

        for prod in products:
            key = prod.model_code.strip().lower()
            if key in existing_rows:
                r_idx = existing_rows[key]
            else:
                current_max += 1
                r_idx = current_max
                existing_rows[key] = r_idx

            sheet.cell(row=r_idx, column=col_indices['model_code'], value=prod.model_code.strip())
            sheet.cell(row=r_idx, column=col_indices['category'], value=prod.category.strip() if prod.category else 'GENERAL')
            sheet.cell(row=r_idx, column=col_indices['description'], value=prod.description.strip() if prod.description else '')

        wb.save(file_path)
        return True, f"Successfully synchronized {products.count()} products to '{filename}'!"

    except PermissionError:
        return False, (
            f"Microsoft Excel currently has '{filename}' open. "
            f"Please close Excel and click 'Sync to Excel' again."
        )
    except Exception as e:
        return False, f"Failed to sync products to Excel: {str(e)}"
