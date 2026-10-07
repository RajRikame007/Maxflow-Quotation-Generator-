from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from .models import Quotation, QuotationItem, Proforma, ProformaItem
from decimal import Decimal
from datetime import date

class QuotationTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user('testadmin', 'admin@example.com', 'pass123')
        self.client.login(username='testadmin', password='pass123')
        self.quotation = Quotation.objects.create(
            quotation_number='TEST-QTN-001',
            quotation_date=date.today(),
            customer_name='Test Client Pvt Ltd',
            customer_address='Test City'
        )
        self.item = QuotationItem.objects.create(
            quotation=self.quotation,
            sr_no='1',
            description='Test Hydraulic Valve',
            quantity=Decimal('2.00'),
            unit='NOS',
            unit_rate=Decimal('5000.00'),
            gst_rate=Decimal('18.00')
        )
        self.quotation.recalculate_totals()

    def test_unauthenticated_user_redirected_to_login(self):
        """Verify that opening the webpage/views before logging in redirects to the login page."""
        anonymous_client = Client()
        
        # Root homepage/dashboard
        res_home = anonymous_client.get(reverse('home'))
        self.assertEqual(res_home.status_code, 302)
        self.assertIn(reverse('login'), res_home.url)
        
        # Quotations history
        res_quotations = anonymous_client.get(reverse('quotation_list'))
        self.assertEqual(res_quotations.status_code, 302)
        self.assertIn(reverse('login'), res_quotations.url)
        
        # Customer master
        res_customers = anonymous_client.get(reverse('customer_list'))
        self.assertEqual(res_customers.status_code, 302)
        self.assertIn(reverse('login'), res_customers.url)
        
        # Product catalog
        res_products = anonymous_client.get(reverse('product_list'))
        self.assertEqual(res_products.status_code, 302)
        self.assertIn(reverse('login'), res_products.url)

    def test_home_view(self):
        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Maxflow Quotation Generator')

    def test_quotation_list_view(self):
        response = self.client.get(reverse('quotation_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'TEST-QTN-001')

    def test_quotation_detail_view(self):
        response = self.client.get(reverse('quotation_detail', args=[self.quotation.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Client Pvt Ltd')

    def test_pdf_download_view(self):
        response = self.client.get(reverse('quotation_pdf_download', args=[self.quotation.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(len(response.content) > 1000)

    def test_login_and_signatory_prefill(self):
        from django.contrib.auth.models import User
        from .models import UserProfile
        user = User.objects.create_user('testuser', 'test@example.com', 'password123', first_name='VIKRAM', last_name='MEHTA')
        profile, _ = UserProfile.objects.get_or_create(user=user)
        profile.designation = 'CHIEF ENGINEER'
        profile.phone = '9876543210'
        profile.save()

        self.client.login(username='testuser', password='password123')
        
        response = self.client.get(reverse('quotation_create'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'VIKRAM MEHTA')
        self.assertContains(response, 'CHIEF ENGINEER')
        self.assertContains(response, '9876543210')

    def test_register_account_with_designation_and_cell_number(self):
        from django.contrib.auth.models import User
        from .models import UserProfile
        self.client.logout()
        
        reg_data = {
            'first_name': 'Rahul',
            'last_name': 'Sharma',
            'username': 'rahul_s',
            'email': 'rahul@maxflowcontrols.com',
            'designation': 'Senior Sales Manager',
            'phone': '9820011223',
            'password1': 'StrongPass123!',
            'password2': 'StrongPass123!',
        }
        response = self.client.post(reverse('register'), data=reg_data, follow=True)
        self.assertEqual(response.status_code, 200)
        
        # Verify user was created
        user = User.objects.get(username='rahul_s')
        self.assertEqual(user.first_name, 'Rahul')
        self.assertEqual(user.last_name, 'Sharma')
        
        # Verify profile was created with designation and cell number
        profile = user.profile
        self.assertEqual(profile.designation, 'SENIOR SALES MANAGER')
        self.assertEqual(profile.phone, '9820011223')
        
        # Verify signatory details directly show in the quotation form
        qtn_response = self.client.get(reverse('quotation_create'))
        self.assertEqual(qtn_response.status_code, 200)
        self.assertContains(qtn_response, 'RAHUL SHARMA')
        self.assertContains(qtn_response, 'SENIOR SALES MANAGER')
        self.assertContains(qtn_response, '9820011223')

    def test_profile_update_view(self):
        from django.contrib.auth.models import User
        user = User.objects.create_user('edituser', 'edit@example.com', 'password123', first_name='Anil', last_name='Patel')
        self.client.login(username='edituser', password='password123')
        
        post_data = {
            'first_name': 'Anil Kumar',
            'last_name': 'Patel',
            'email': 'anil.patel@maxflow.com',
            'designation': 'Technical Director',
            'phone': '9123456789'
        }
        response = self.client.post(reverse('profile'), data=post_data, follow=True)
        self.assertEqual(response.status_code, 200)
        
        user.refresh_from_db()
        self.assertEqual(user.first_name, 'Anil Kumar')
        self.assertEqual(user.profile.designation, 'TECHNICAL DIRECTOR')
        self.assertEqual(user.profile.phone, '9123456789')

    def test_customer_list_and_search(self):
        from .models import Customer
        Customer.objects.create(
            name='JSW STEEL COATED PRODUCTS LIMITED',
            mobile='8709305772',
            gstin='27AACCM3988L1ZU',
            address='VILLAGE - VASIND, TALUKA - SHAHAPUR'
        )
        response = self.client.get(reverse('customer_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'JSW STEEL COATED PRODUCTS LIMITED')

        # Test search
        search_res = self.client.get(reverse('customer_list') + '?q=VASIND')
        self.assertContains(search_res, 'JSW STEEL COATED PRODUCTS LIMITED')

    def test_customer_search_api(self):
        from .models import Customer
        Customer.objects.create(
            name='ARV ENGINEERING CO. LLP',
            email='purchase1@arvengg.com',
            gstin='27ABJFA2735N1Z4'
        )
        response = self.client.get(reverse('customer_search_api') + '?q=ARV')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(len(data['customers']) > 0)
        self.assertEqual(data['customers'][0]['name'], 'ARV ENGINEERING CO. LLP')

    def test_product_search_api_and_list(self):
        from .models import Product
        Product.objects.create(
            category='EATON',
            model_code='6033556-001',
            description='24V DC COIL (EN124)',
            unit_rate=Decimal('2500.00'),
            is_active=True
        )
        Product.objects.create(
            category='HYDROLINE',
            model_code='003-SC3-020',
            description='STRAINER',
            unit_rate=Decimal('1200.00'),
            is_active=True
        )

        # 1. Search by code
        res1 = self.client.get(reverse('product_search_api') + '?q=6033556')
        self.assertEqual(res1.status_code, 200)
        data1 = res1.json()
        self.assertTrue(len(data1['products']) >= 1)
        self.assertEqual(data1['products'][0]['code'], '6033556-001')
        self.assertIn('24V DC COIL', data1['products'][0]['display_desc'])

        # 2. Search by flexible code ignoring hyphens, spaces, and letter case (e.g. v21051a)
        Product.objects.create(
            category='EATON',
            model_code='V210-5-1A-12-S214-INB',
            description='VANE PUMP',
            unit_rate=Decimal('12000.00'),
            is_active=True
        )
        res_v1 = self.client.get(reverse('product_search_api') + '?q=v21051a')
        self.assertEqual(res_v1.status_code, 200)
        v_codes = [p['code'] for p in res_v1.json()['products']]
        self.assertIn('V210-5-1A-12-S214-INB', v_codes)

        res_v2 = self.client.get(reverse('product_search_api') + '?q=v210 5 1a')
        self.assertEqual(res_v2.status_code, 200)
        self.assertIn('V210-5-1A-12-S214-INB', [p['code'] for p in res_v2.json()['products']])

        # 3. Search by brand
        res2 = self.client.get(reverse('product_search_api') + '?brand=HYDROLINE')
        self.assertEqual(res2.status_code, 200)
        data2 = res2.json()
        self.assertTrue(any(p['code'] == '003-SC3-020' for p in data2['products']))

        # 4. Product directory list page with flexible query
        res3 = self.client.get(reverse('product_list') + '?q=v21051a')
        self.assertEqual(res3.status_code, 200)
        self.assertContains(res3, 'V210-5-1A-12-S214-INB')

    def test_quotation_create_with_customer_id(self):
        from .models import Customer
        cust = Customer.objects.create(
            name='MAHINDRA & MAHINDRA LTD.',
            mobile='9821356433',
            gstin='27AAACM3025E1ZZ',
            address='Kandivli (E), Mumbai'
        )
        response = self.client.get(reverse('quotation_create') + f'?customer_id={cust.id}')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'MAHINDRA &amp; MAHINDRA LTD.')
        self.assertContains(response, '27AAACM3025E1ZZ')

    def test_customer_create_post(self):
        post_data = {
            'name': 'ZENITH VALVES INDIA',
            'gstin': '27AAACZ1234F1Z9',
            'mobile': '9820011223',
            'email': 'sales@zenithvalves.com',
            'address': 'Plot 45, Rabale MIDC, Navi Mumbai'
        }
        response = self.client.post(reverse('customer_create'), data=post_data, follow=True)
        self.assertEqual(response.status_code, 200)
        from .models import Customer
        self.assertTrue(Customer.objects.filter(name='ZENITH VALVES INDIA').exists())

    def test_customer_create_ajax(self):
        post_data = {
            'name': 'APEX HYDRAULICS LTD',
            'gstin': '27AAACA9999M1Z1',
            'mobile': '9876543210',
            'email': 'info@apex.com',
            'address': 'Thane West'
        }
        response = self.client.post(
            reverse('customer_create'),
            data=post_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['customer']['name'], 'APEX HYDRAULICS LTD')

    def test_exact_pf_freight_and_gst_calculation(self):
        q = Quotation.objects.create(
            quotation_number='TEST-CALC-PERFECT',
            quotation_date=date.today(),
            customer_name='CALC TEST CLIENT',
            discount_percentage=Decimal('5.00'),
            pf_percentage=Decimal('2.50'),
            freight_amount=Decimal('350.00')
        )
        QuotationItem.objects.create(
            quotation=q,
            sr_no='1',
            description='Hydraulic Pump Model A',
            quantity=Decimal('5.00'),
            unit='NOS',
            unit_rate=Decimal('1234.56'),
            gst_rate=Decimal('18.00')
        )
        q.recalculate_totals()

        self.assertEqual(q.subtotal, Decimal('6172.80'))
        self.assertEqual(q.discount_amount, Decimal('308.64'))
        self.assertEqual(q.pf_amount, Decimal('146.60'))
        self.assertEqual(q.freight_amount, Decimal('350.00'))
        self.assertEqual(q.taxable_amount, Decimal('6360.76'))
        self.assertEqual(q.tax_amount, Decimal('1144.94'))
        self.assertEqual(q.grand_total, Decimal('7505.70'))

    def test_quotation_edit_with_blank_pf_percentage(self):
        """Ensure submitting a quotation edit with blank/empty pf_percentage succeeds without NOT NULL error."""
        post_data = {
            'company_name': self.quotation.company_name,
            'company_address': self.quotation.company_address,
            'company_phone': self.quotation.company_phone,
            'company_email': self.quotation.company_email,
            'quotation_number': self.quotation.quotation_number,
            'quotation_date': str(self.quotation.quotation_date),
            'salutation': 'DEAR SIR,',
            'subject': 'QUOTATION NOTE',
            'customer_name': self.quotation.customer_name,
            'customer_address': self.quotation.customer_address,
            'price_terms': 'F.O.R., DESTINATION',
            'freight_terms': 'EXTRA',
            'pf_terms': 'NIL',
            'discount_terms': 'NET.',
            'tax_terms': '18% GST',
            'payment_terms': '30 DAYS',
            'validity_terms': '30 DAYS.',
            'signatory_company': 'MAXFLOW',
            'signatory_name': 'DINENDRA CHARI',
            'signatory_designation': 'MANAGER',
            # Blank/empty financial fields to reproduce the issue
            'discount_percentage': '',
            'discount_amount': '',
            'pf_percentage': '',
            'pf_amount': '',
            'freight_amount': '',
            'taxable_amount': '',
            'subtotal': '',
            'tax_amount': '',
            'grand_total': '',
            # Formset management form
            'items-TOTAL_FORMS': '1',
            'items-INITIAL_FORMS': '1',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-id': str(self.item.id),
            'items-0-sr_no': '1',
            'items-0-description': 'Testing Item Without PF',
            'items-0-quantity': '1',
            'items-0-unit': 'NOS',
            'items-0-unit_rate': '1000.00',
            'items-0-gst_rate': '18.00',
            'items-0-amount': '1000.00',
        }
        response = self.client.post(reverse('quotation_edit', args=[self.quotation.pk]), data=post_data)
        # Should redirect to quotation_detail on success
        self.assertEqual(response.status_code, 302)
        self.quotation.refresh_from_db()
        self.assertEqual(self.quotation.pf_percentage, Decimal('0.00'))
        self.assertEqual(self.quotation.subtotal, Decimal('1000.00'))

    def test_quotation_edit_custom_sr_no_preserved(self):
        """Ensure custom item serial numbers (e.g. 'Item 1A', '2.1') are preserved on edit and shown in detail and PDF."""
        post_data = {
            'company_name': self.quotation.company_name,
            'company_address': self.quotation.company_address,
            'company_phone': self.quotation.company_phone,
            'company_email': self.quotation.company_email,
            'quotation_number': self.quotation.quotation_number,
            'quotation_date': str(self.quotation.quotation_date),
            'salutation': 'DEAR SIR,',
            'subject': 'QUOTATION NOTE',
            'customer_name': self.quotation.customer_name,
            'customer_address': self.quotation.customer_address,
            'price_terms': 'F.O.R., DESTINATION',
            'freight_terms': 'EXTRA',
            'pf_terms': 'NIL',
            'discount_terms': 'NET.',
            'tax_terms': '18% GST Extra',
            'payment_terms': '30 DAYS',
            'validity_terms': '30 DAYS.',
            'signatory_company': 'MAXFLOW',
            'signatory_name': 'DINENDRA CHARI',
            'signatory_designation': 'MANAGER',
            'items-TOTAL_FORMS': '2',
            'items-INITIAL_FORMS': '1',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-id': str(self.item.id),
            'items-0-sr_no': 'Item 1A',
            'items-0-description': 'First Item Custom SR',
            'items-0-quantity': '1',
            'items-0-unit': 'NOS',
            'items-0-unit_rate': '1000.00',
            'items-0-gst_rate': '18',
            'items-0-amount': '1000.00',
            'items-1-id': '',
            'items-1-sr_no': 'Item 2B',
            'items-1-description': 'Second Item Custom SR',
            'items-1-quantity': '2',
            'items-1-unit': 'NOS',
            'items-1-unit_rate': '2000.00',
            'items-1-gst_rate': '18%',
            'items-1-amount': '4000.00',
        }
        response = self.client.post(reverse('quotation_edit', args=[self.quotation.pk]), data=post_data)
        self.assertEqual(response.status_code, 302)

        self.item.refresh_from_db()
        self.assertEqual(self.item.sr_no, 'Item 1A')

        new_item = self.quotation.items.exclude(pk=self.item.pk).first()
        self.assertIsNotNone(new_item)
        self.assertEqual(new_item.sr_no, 'Item 2B')

        # Check detail view displays custom sr_no
        detail_resp = self.client.get(reverse('quotation_detail', args=[self.quotation.pk]))
        self.assertContains(detail_resp, 'Item 1A')
        self.assertContains(detail_resp, 'Item 2B')

    def test_quotation_edit_add_gst_after_saving_without_gst(self):
        """Ensure editing a quotation to add GST after it was saved without GST calculates and persists GST correctly."""
        # Create a quotation without GST
        no_gst_q = Quotation.objects.create(
            quotation_number='TEST-NOGST-99',
            quotation_date=date.today(),
            customer_name='Zero GST Client',
            customer_address='Zero Tax Lane',
            tax_terms='GST Extra'
        )
        q_item = QuotationItem.objects.create(
            quotation=no_gst_q,
            sr_no='1',
            description='Machine Part Without Initial Tax',
            quantity=Decimal('2.00'),
            unit='NOS',
            unit_rate=Decimal('5000.00'),
            gst_rate=None
        )
        no_gst_q.recalculate_totals()
        self.assertEqual(no_gst_q.tax_amount, Decimal('0.00'))
        self.assertEqual(no_gst_q.grand_total, Decimal('10000.00'))

        # Now edit to add 18% GST (testing with '18%')
        post_data = {
            'company_name': no_gst_q.company_name,
            'company_address': no_gst_q.company_address,
            'company_phone': no_gst_q.company_phone,
            'company_email': no_gst_q.company_email,
            'quotation_number': no_gst_q.quotation_number,
            'quotation_date': str(no_gst_q.quotation_date),
            'salutation': 'DEAR SIR,',
            'subject': 'QUOTATION NOTE',
            'customer_name': no_gst_q.customer_name,
            'customer_address': no_gst_q.customer_address,
            'price_terms': 'F.O.R., DESTINATION',
            'freight_terms': 'EXTRA',
            'pf_terms': 'NIL',
            'discount_terms': 'NET.',
            'tax_terms': 'AS INDICATED ABOVE (18% GST Extra)',
            'payment_terms': '30 DAYS',
            'validity_terms': '30 DAYS.',
            'signatory_company': 'MAXFLOW',
            'signatory_name': 'DINENDRA CHARI',
            'signatory_designation': 'MANAGER',
            'items-TOTAL_FORMS': '1',
            'items-INITIAL_FORMS': '1',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-id': str(q_item.id),
            'items-0-sr_no': '1',
            'items-0-description': 'Machine Part Without Initial Tax',
            'items-0-quantity': '2',
            'items-0-unit': 'NOS',
            'items-0-unit_rate': '5000.00',
            'items-0-gst_rate': '18%',
            'items-0-amount': '10000.00',
        }
        edit_resp = self.client.post(reverse('quotation_edit', args=[no_gst_q.pk]), data=post_data)
        self.assertEqual(edit_resp.status_code, 302)

        no_gst_q.refresh_from_db()
        q_item.refresh_from_db()
        self.assertEqual(q_item.gst_rate, Decimal('18.00'))
        self.assertEqual(no_gst_q.tax_amount, Decimal('1800.00'))
        self.assertEqual(no_gst_q.grand_total, Decimal('11800.00'))
        self.assertIn('18%', no_gst_q.gst_rates_display)

        # PDF download should generate successfully and have 18% tax
        pdf_resp = self.client.get(reverse('quotation_pdf_download', args=[no_gst_q.pk]))
        self.assertEqual(pdf_resp.status_code, 200)
        self.assertEqual(pdf_resp['Content-Type'], 'application/pdf')

        no_gst_q.delete()

    def test_customer_edit_post(self):
        from .models import Customer
        cust = Customer.objects.create(
            name='DELTA ENGINEERING WORKS',
            mobile='9822001122',
            gstin='27AAACD1111E1Z0',
            address='Old Address, Pune'
        )
        post_data = {
            'name': 'DELTA ENGINEERING WORKS PVT LTD',
            'mobile': '9822998877',
            'email': 'contact@deltaengg.com',
            'gstin': '27AAACD1111E1Z0',
            'address': 'New Modern Facility, Hinjawadi Phase 2, Pune'
        }
        response = self.client.post(reverse('customer_edit', args=[cust.id]), data=post_data, follow=True)
        self.assertEqual(response.status_code, 200)

        cust.refresh_from_db()
        self.assertEqual(cust.name, 'DELTA ENGINEERING WORKS PVT LTD')
        self.assertEqual(cust.mobile, '9822998877')
        self.assertEqual(cust.email, 'contact@deltaengg.com')
        self.assertIn('Hinjawadi', cust.address)

    def test_customer_edit_ajax(self):
        from .models import Customer
        cust = Customer.objects.create(
            name='ALPHA PRECISION TOOLS',
            mobile='9123456789',
            gstin='27AAACA1234F1Z1',
            address='Nashik'
        )
        post_data = {
            'name': 'ALPHA PRECISION TOOLS & DIES',
            'mobile': '9988776655',
            'email': 'alpha@precision.com',
            'gstin': '27AAACA1234F1Z1',
            'address': 'Ambad MIDC, Nashik - 422010'
        }
        response = self.client.post(
            reverse('customer_edit', args=[cust.id]),
            data=post_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertEqual(data['customer']['name'], 'ALPHA PRECISION TOOLS & DIES')
        self.assertEqual(data['customer']['mobile'], '9988776655')

        cust.refresh_from_db()
        self.assertEqual(cust.name, 'ALPHA PRECISION TOOLS & DIES')

    def test_customer_edit_excel_sync(self):
        import tempfile
        import openpyxl
        from .models import Customer
        from .excel_sync import append_or_update_customer_in_excel

        # Create temporary workbook simulating customer address raj.xlsx
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value='PARTY NAME')
        ws.cell(row=2, column=2, value='EMAIL')
        ws.cell(row=2, column=3, value='MOBILE')
        ws.cell(row=2, column=4, value='PARTY GSTN')
        ws.cell(row=2, column=5, value='FULL ADDRESS IN TALLY')

        # Add initial row
        ws.cell(row=3, column=1, value='OMEGA VALVES')
        ws.cell(row=3, column=2, value='omega@old.com')
        ws.cell(row=3, column=3, value='9000000000')
        ws.cell(row=3, column=4, value='27AAACT9999M1Z9')
        ws.cell(row=3, column=5, value='Old Plant, Rabale')

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name
        wb.save(tmp_path)

        try:
            cust = Customer.objects.create(
                name='OMEGA VALVES INTERNATIONAL',
                email='sales@omegavalves.com',
                mobile='9111111111',
                gstin='27AAACT9999M1Z9',
                address='New Global Center, Navi Mumbai'
            )

            success, msg = append_or_update_customer_in_excel(
                cust,
                original_name='OMEGA VALVES',
                filename=tmp_path
            )
            self.assertTrue(success)

            # Verify the same row 3 was updated rather than appending row 4
            updated_wb = openpyxl.load_workbook(tmp_path)
            updated_ws = updated_wb.active
            self.assertEqual(updated_ws.max_row, 3)
            self.assertEqual(updated_ws.cell(row=3, column=1).value, 'OMEGA VALVES INTERNATIONAL')
            self.assertEqual(updated_ws.cell(row=3, column=2).value, 'sales@omegavalves.com')
            self.assertEqual(updated_ws.cell(row=3, column=3).value, '9111111111')
            self.assertEqual(updated_ws.cell(row=3, column=5).value, 'New Global Center, Navi Mumbai')
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_customer_delete_and_excel_sync(self):
        import tempfile
        import openpyxl
        from .models import Customer
        from .excel_sync import delete_customer_from_excel

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value='PARTY NAME')
        ws.cell(row=2, column=2, value='EMAIL')
        ws.cell(row=2, column=3, value='MOBILE')
        ws.cell(row=2, column=4, value='PARTY GSTN')
        ws.cell(row=2, column=5, value='FULL ADDRESS IN TALLY')

        ws.cell(row=3, column=1, value='ALPHA MOTORS PVT LTD')
        ws.cell(row=3, column=4, value='27AAACA1234A1Z1')
        ws.cell(row=4, column=1, value='BETA POWER LTD')
        ws.cell(row=4, column=4, value='27AAACB5678B1Z2')

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name
        wb.save(tmp_path)

        try:
            cust = Customer.objects.create(
                name='ALPHA MOTORS PVT LTD',
                gstin='27AAACA1234A1Z1'
            )
            success, msg = delete_customer_from_excel(
                cust.name,
                gstin=cust.gstin,
                filename=tmp_path
            )
            self.assertTrue(success)

            updated_wb = openpyxl.load_workbook(tmp_path)
            updated_ws = updated_wb.active
            # After deleting row 3, row 4 (BETA POWER LTD) should now be row 3
            self.assertEqual(updated_ws.cell(row=3, column=1).value, 'BETA POWER LTD')
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_product_create_and_excel_sync(self):
        import tempfile
        import openpyxl
        from .models import Product
        from .excel_sync import append_or_update_product_in_excel

        # Create temporary workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value='Stock item Name')
        ws.cell(row=2, column=2, value='BRAND')
        ws.cell(row=2, column=3, value='DESCRIPTION')

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name
        wb.save(tmp_path)

        try:
            prod = Product.objects.create(
                category='DANFOSS',
                model_code='OMM-32-151G0001',
                description='HYDRAULIC MOTOR 32CC',
                unit_rate=Decimal('14500.00'),
                is_active=True
            )
            success, msg = append_or_update_product_in_excel(prod, filename=tmp_path)
            self.assertTrue(success)

            updated_wb = openpyxl.load_workbook(tmp_path)
            updated_ws = updated_wb.active
            self.assertEqual(updated_ws.max_row, 3)
            self.assertEqual(updated_ws.cell(row=3, column=1).value, 'OMM-32-151G0001')
            self.assertEqual(updated_ws.cell(row=3, column=2).value, 'DANFOSS')
            self.assertEqual(updated_ws.cell(row=3, column=3).value, 'HYDRAULIC MOTOR 32CC')
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_product_edit_and_excel_sync(self):
        import tempfile
        import openpyxl
        from .models import Product
        from .excel_sync import append_or_update_product_in_excel

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value='Stock item Name')
        ws.cell(row=2, column=2, value='BRAND')
        ws.cell(row=2, column=3, value='DESCRIPTION')

        ws.cell(row=3, column=1, value='OLD-CODE-100')
        ws.cell(row=3, column=2, value='EATON')
        ws.cell(row=3, column=3, value='OLD DESCRIPTION')

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name
        wb.save(tmp_path)

        try:
            prod = Product.objects.create(
                category='EATON VICKERS',
                model_code='NEW-CODE-200',
                description='UPDATED VALVE 24V',
                unit_rate=Decimal('8200.00'),
                is_active=True
            )
            # Edit matching row 3 by original_code
            success, msg = append_or_update_product_in_excel(
                prod,
                original_code='OLD-CODE-100',
                filename=tmp_path
            )
            self.assertTrue(success)

            updated_wb = openpyxl.load_workbook(tmp_path)
            updated_ws = updated_wb.active
            self.assertEqual(updated_ws.max_row, 3) # Same row updated
            self.assertEqual(updated_ws.cell(row=3, column=1).value, 'NEW-CODE-200')
            self.assertEqual(updated_ws.cell(row=3, column=2).value, 'EATON VICKERS')
            self.assertEqual(updated_ws.cell(row=3, column=3).value, 'UPDATED VALVE 24V')
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_product_delete_and_excel_sync(self):
        import tempfile
        import openpyxl
        from .excel_sync import delete_product_from_excel

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value='Stock item Name')
        ws.cell(row=2, column=2, value='BRAND')
        ws.cell(row=2, column=3, value='DESCRIPTION')

        ws.cell(row=3, column=1, value='PROD-TO-DELETE')
        ws.cell(row=3, column=2, value='HYDROLINE')
        ws.cell(row=3, column=3, value='FILTER ELEMENT')

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name
        wb.save(tmp_path)

        try:
            success, msg = delete_product_from_excel('PROD-TO-DELETE', category='HYDROLINE', filename=tmp_path)
            self.assertTrue(success)

            updated_wb = openpyxl.load_workbook(tmp_path)
            updated_ws = updated_wb.active
            self.assertEqual(updated_ws.max_row, 2) # Row 3 deleted
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_product_views_ajax_crud(self):
        from .models import Product

        # 1. Add Product via AJAX
        add_data = {
            'model_code': 'AJAX-VALVE-01',
            'category': 'YUKEN',
            'description': 'DIRECTIONAL VALVE DSG-01-3C2',
            'unit_rate': '4500.00',
            'hsn_code': '84812000',
            'unit': 'NOS',
            'gst_rate': '18.00'
        }
        res_add = self.client.post(
            reverse('product_create'),
            data=add_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(res_add.status_code, 200)
        data_add = res_add.json()
        self.assertTrue(data_add['success'])
        prod_id = data_add['product']['id']

        # Verify DB
        prod = Product.objects.get(id=prod_id)
        self.assertEqual(prod.model_code, 'AJAX-VALVE-01')

        # 2. Edit Product via AJAX
        edit_data = {
            'model_code': 'AJAX-VALVE-01-MOD',
            'category': 'YUKEN KOGYO',
            'description': 'MODIFIED VALVE DSG-01-3C2-A240',
            'unit_rate': '5200.00',
            'hsn_code': '84812000',
            'unit': 'NOS',
            'gst_rate': '18.00'
        }
        res_edit = self.client.post(
            reverse('product_edit', args=[prod_id]),
            data=edit_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(res_edit.status_code, 200)
        data_edit = res_edit.json()
        self.assertTrue(data_edit['success'])
        self.assertEqual(data_edit['product']['model_code'], 'AJAX-VALVE-01-MOD')

        prod.refresh_from_db()
        self.assertEqual(prod.model_code, 'AJAX-VALVE-01-MOD')

        # 3. Delete Product via AJAX
        res_del = self.client.post(
            reverse('product_delete', args=[prod_id]),
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(res_del.status_code, 200)
        data_del = res_del.json()
        self.assertTrue(data_del['success'])
        self.assertFalse(Product.objects.filter(id=prod_id).exists())

    def test_customer_quick_save_api(self):
        from .models import Customer
        from django.urls import reverse

        # 1. Quick save new customer
        new_data = {
            'name': 'TEST NEW MANUAL CORP',
            'address': 'Plot 45, Phase 2, GIDC, Naroda, Ahmedabad',
            'phone': '9876543210',
            'email': 'contact@manualcorp.com',
            'gstin': '24AAACT1234M1Z5'
        }
        res = self.client.post(
            reverse('customer_quick_save'),
            data=new_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(res.status_code, 200)
        json_data = res.json()
        self.assertTrue(json_data['success'])
        self.assertTrue(json_data['is_new'])
        self.assertEqual(json_data['customer']['name'], 'TEST NEW MANUAL CORP')

        cust = Customer.objects.get(name='TEST NEW MANUAL CORP')
        self.assertEqual(cust.mobile, '9876543210')

        # 2. Quick update existing customer
        update_data = {
            'name': 'TEST NEW MANUAL CORP',
            'address': 'Updated Address Street 10',
            'phone': '9999988888',
            'email': 'updated@manualcorp.com',
            'gstin': '24AAACT1234M1Z5'
        }
        res2 = self.client.post(
            reverse('customer_quick_save'),
            data=update_data,
            HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )
        self.assertEqual(res2.status_code, 200)
        json_data2 = res2.json()
        self.assertTrue(json_data2['success'])
        self.assertFalse(json_data2['is_new'])

        cust.refresh_from_db()
        self.assertEqual(cust.address, 'Updated Address Street 10')
        self.assertEqual(cust.mobile, '9999988888')

    def test_quotation_save_customer_to_master(self):
        from .models import Quotation, Customer
        from django.urls import reverse

        post_data = {
            'company_name': 'Maxflow Controls ( I ) Pvt. Ltd.',
            'company_address': '18, Unique Industrial Estate, Prabhadevi, Mumbai',
            'company_phone': '022-2436 0131',
            'company_email': 'mumbai@maxflowcontrols.com',
            'quotation_number': 'QTN/TEST/AUTO/001',
            'quotation_date': '2026-10-02',
            'salutation': 'DEAR SIR,',
            'subject': 'QUOTATION SUBJECT',
            'customer_name': 'AUTO SYNC CLIENT PVT LTD',
            'customer_address': 'Plot 99, MIDC Taloja, Navi Mumbai',
            'customer_phone': '022-27412345',
            'customer_email': 'sales@autosync.com',
            'customer_gstin': '27AAACA9999Z1Z0',
            'price_terms': 'F.O.R., DESTINATION',
            'freight_terms': 'EXTRA',
            'pf_terms': 'NIL',
            'delivery_terms': '4 WEEKS',
            'payment_terms': '30 DAYS',
            'validity_terms': '30 DAYS',
            'tax_terms': 'GST 18% EXTRA',
            'discount_terms': 'NIL',
            'warranty_terms': '12 MONTHS',
            'signatory_company': 'Maxflow Controls',
            'signatory_name': 'TEST USER',
            'signatory_designation': 'MANAGER',
            'save_customer_to_master': 'on',
            'items-TOTAL_FORMS': '1',
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-sr_no': '1',
            'items-0-description': 'TEST HYDRAULIC MOTOR',
            'items-0-hsn_code': '84136090',
            'items-0-quantity': '1',
            'items-0-unit': 'NOS',
            'items-0-unit_rate': '10000.00',
            'items-0-delivery_schedule': 'EX STOCK',
            'items-0-gst_rate': '18.00',
            'items-0-amount': '10000.00',
        }

        res = self.client.post(reverse('quotation_create'), data=post_data)
        self.assertEqual(res.status_code, 302)

        # Customer must have been created in Master database
        cust = Customer.objects.filter(name='AUTO SYNC CLIENT PVT LTD').first()
        self.assertIsNotNone(cust)
        self.assertEqual(cust.mobile, '022-27412345')
        self.assertEqual(cust.gstin, '27AAACA9999Z1Z0')


class ProformaTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user('adminproforma', 'proforma@example.com', 'secret123')
        self.client.login(username='adminproforma', password='secret123')

        # Create source quotation
        self.quotation = Quotation.objects.create(
            quotation_number='QTN-SOURCE-001',
            quotation_date=date.today(),
            customer_name='HydraSpares & Engineering.',
            customer_address='Aj-65/3, Khudiram Pally, Talpukur Road, Kolkata - 700061',
            customer_phone='9804736661',
            customer_gstin='19DTYPG1669F1Z1',
            freight_amount=Decimal('550.00'),
        )
        self.item1 = QuotationItem.objects.create(
            quotation=self.quotation,
            sr_no='1',
            description='Cartridge Valve – Part# 406AA00066A\nModel# 1CEB120P35S3',
            quantity=Decimal('1.00'),
            unit='NOS',
            unit_rate=Decimal('12534.00'),
            gst_rate=Decimal('18.00')
        )
        self.quotation.recalculate_totals()

    def test_generate_proforma_from_quotation(self):
        """Test clicking Generate Proforma copies data and leaves quotation untouched."""
        orig_qtn_total = self.quotation.grand_total

        res = self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))
        self.assertEqual(res.status_code, 302)

        # Proforma should be created
        proforma = Proforma.objects.filter(quotation=self.quotation).first()
        self.assertIsNotNone(proforma)
        self.assertEqual(proforma.customer_name, self.quotation.customer_name)
        self.assertEqual(proforma.customer_gstin, self.quotation.customer_gstin)
        self.assertEqual(proforma.items.count(), 1)
        self.assertEqual(proforma.freight_amount, Decimal('0.00'))
        self.assertEqual(proforma.items.first().unit_rate, Decimal('12534.00'))

        # Verify source quotation is completely unmodified
        self.quotation.refresh_from_db()
        self.assertEqual(self.quotation.grand_total, orig_qtn_total)

    def test_edit_proforma_does_not_modify_original_quotation(self):
        """Editing Proforma items and totals must NOT alter source quotation."""
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))
        proforma = Proforma.objects.filter(quotation=self.quotation).first()
        orig_qtn_total = self.quotation.grand_total

        # Edit Proforma item
        p_item = proforma.items.first()
        p_item.quantity = Decimal('8.00')
        p_item.unit_rate = Decimal('10000.00')
        p_item.save()
        proforma.recalculate_totals()

        # Proforma total changed
        self.assertNotEqual(proforma.grand_total, orig_qtn_total)

        # Original quotation unchanged
        self.quotation.refresh_from_db()
        self.assertEqual(self.quotation.grand_total, orig_qtn_total)
        self.assertEqual(self.quotation.items.first().quantity, Decimal('1.00'))

    def test_multiple_proformas_from_single_quotation(self):
        """A single quotation can generate multiple distinct Proforma Invoices."""
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))

        proformas = Proforma.objects.filter(quotation=self.quotation)
        self.assertEqual(proformas.count(), 2)
        p1, p2 = proformas[0], proformas[1]
        self.assertNotEqual(p1.proforma_number, p2.proforma_number)

    def test_proforma_pdf_preview_and_download(self):
        """Verify PDF download and preview return valid PDF stream."""
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))
        proforma = Proforma.objects.filter(quotation=self.quotation).first()

        # Download view
        dl_res = self.client.get(reverse('proforma_pdf_download', args=[proforma.pk]))
        self.assertEqual(dl_res.status_code, 200)
        self.assertEqual(dl_res['Content-Type'], 'application/pdf')
        self.assertIn('attachment', dl_res['Content-Disposition'])
        self.assertTrue(len(dl_res.content) > 1000)

        # Inline preview
        preview_res = self.client.get(reverse('proforma_pdf_preview', args=[proforma.pk]))
        self.assertEqual(preview_res.status_code, 200)
        self.assertEqual(preview_res['Content-Type'], 'application/pdf')
        self.assertIn('inline', preview_res['Content-Disposition'])

    def test_proforma_history_and_dashboard(self):
        """Verify Proforma History list and Dashboard statistics."""
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))

        # History list
        list_res = self.client.get(reverse('proforma_list'))
        self.assertEqual(list_res.status_code, 200)
        self.assertContains(list_res, 'HydraSpares')

        # Dashboard
        dash_res = self.client.get(reverse('home'))
        self.assertEqual(dash_res.status_code, 200)
        self.assertContains(dash_res, 'Total Proformas')
        self.assertContains(dash_res, 'Proformas This Month')

    def test_proforma_full_calculation_with_discount_rounding_and_advance(self):
        """
        Verify the exact calculation from user's specification & screenshot:
        1. Item: 2 * 1,60,275.00 = 3,20,550.00
        2. Discount 15% LESS = 48,082.50
        3. Subtotal after discount = 2,72,467.50
        4. GST @18% = 49,044.15
        5. Total = 3,21,511.65
        6. Rounding Off Total = 3,21,512.00
        7. Advance Received = 86,000.00
        8. Balance Amount Payable = 2,35,512.00
        """
        proforma = Proforma.objects.create(
            proforma_number='MCIPL/03/2627/0001/DAN',
            proforma_date=date(2026, 3, 13),
            customer_name='DANFOSS CLIENT',
            customer_address='Pune, Maharashtra',
            tax_type='IGST',
            tax_rate=Decimal('18.00'),
            discount_percentage=Decimal('15.00'),
            round_off_enabled=True,
            advance_label='ADVANCE RECEIVED ON OUR A/C',
            advance_date=date(2026, 3, 13),
            advance_amount=Decimal('86000.00')
        )
        ProformaItem.objects.create(
            proforma=proforma,
            sr_no='1.',
            description='DANFOSS VICKERS MAKE PISTON PUMP MODEL : PVM018ER05CS1C28011000AAB-0000',
            quantity=Decimal('2.00'),
            unit='No',
            unit_rate=Decimal('160275.00')
        )

        proforma.recalculate_totals()

        self.assertEqual(proforma.subtotal, Decimal('320550.00'))
        self.assertEqual(proforma.discount_amount, Decimal('48082.50'))
        self.assertEqual(proforma.subtotal_after_discount, Decimal('272467.50'))
        self.assertEqual(proforma.taxable_amount, Decimal('272467.50'))
        self.assertEqual(proforma.igst_amount, Decimal('49044.15'))
        self.assertEqual(proforma.grand_total, Decimal('321511.65'))
        self.assertEqual(proforma.rounded_total, Decimal('321512.00'))
        self.assertEqual(proforma.round_off_amount, Decimal('0.35'))
        self.assertEqual(proforma.advance_amount, Decimal('86000.00'))
        self.assertEqual(proforma.balance_payable, Decimal('235512.00'))
        self.assertIn('TWO LAKH THIRTY FIVE THOUSAND FIVE HUNDRED TWELVE', proforma.amount_in_words.upper())

        # Verify PDF renders these rows
        pdf_res = self.client.get(reverse('proforma_pdf_preview', args=[proforma.pk]))
        self.assertEqual(pdf_res.status_code, 200)

        # Verify detail page renders these rows
        detail_res = self.client.get(reverse('proforma_detail', args=[proforma.pk]))
        self.assertEqual(detail_res.status_code, 200)
        self.assertContains(detail_res, 'DISCOUNT 15% LESS')
        self.assertContains(detail_res, '48,082.50')
        self.assertContains(detail_res, '272,467.50')
        self.assertContains(detail_res, 'GST ADD @18%')
        self.assertContains(detail_res, '49,044.15')
        self.assertContains(detail_res, '321,511.65')
        self.assertContains(detail_res, 'ROUNDING OFF TOTAL')
        self.assertContains(detail_res, '321,512.00')
        self.assertContains(detail_res, 'ADVANCE RECEIVED ON OUR A/C DT 13.03.26')
        self.assertContains(detail_res, '86,000.00')
        self.assertContains(detail_res, 'BALANCE AMOUNT PAYABLE BY YOU')
        self.assertContains(detail_res, '235,512.00')

    def test_proforma_edit_with_advance_and_round_off(self):
        """Test editing a proforma through proforma_edit view saves advance and round off correctly."""
        self.client.get(reverse('proforma_generate_from_quotation', args=[self.quotation.pk]))
        proforma = Proforma.objects.filter(quotation=self.quotation).first()
        edit_url = reverse('proforma_edit', args=[proforma.pk])

        post_data = {
            'proforma_number': proforma.proforma_number,
            'proforma_date': '2026-10-03',
            'status': 'Draft',
            'company_name': proforma.company_name,
            'company_address': proforma.company_address,
            'company_contact': proforma.company_contact,
            'customer_name': proforma.customer_name,
            'customer_address': proforma.customer_address,
            'customer_contact': '',
            'customer_gstin': '',
            'customer_email': '',
            'attention_to': '',
            'salutation': proforma.salutation,
            'po_reference': proforma.po_reference,
            'po_date': '2026-10-03',
            'subject_note': proforma.subject_note,
            'freight_label': 'DTDC BY AIR',
            'freight_amount': '0.00',
            'tax_type': 'GST',
            'tax_rate': '18.00',
            'discount_percentage': '15.00',
            'round_off_enabled': 'on',
            'advance_label': 'ADVANCE RECEIVED ON OUR A/C',
            'advance_date': '2026-03-13',
            'advance_amount': '86000.00',
            'bank_name': proforma.bank_name,
            'bank_branch': proforma.bank_branch,
            'bank_address': proforma.bank_address,
            'bank_telephone': proforma.bank_telephone,
            'bank_account_name': proforma.bank_account_name,
            'bank_account_no': proforma.bank_account_no,
            'bank_ifsc': proforma.bank_ifsc,
            'bank_micr': proforma.bank_micr,
            'request_note': proforma.request_note,
            'signatory_company': proforma.signatory_company,
            'signatory_name': proforma.signatory_name,
            'signatory_designation': proforma.signatory_designation,
            # Formset management form
            'items-TOTAL_FORMS': '1',
            'items-INITIAL_FORMS': '1',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-id': str(proforma.items.first().pk),
            'items-0-sr_no': '1.',
            'items-0-description': 'Danfoss Piston Pump',
            'items-0-quantity': '2',
            'items-0-unit': 'No',
            'items-0-unit_rate': '160275.00',
        }

        response = self.client.post(edit_url, post_data)
        self.assertEqual(response.status_code, 302)

        proforma.refresh_from_db()
        self.assertEqual(proforma.subtotal, Decimal('320550.00'))
        self.assertEqual(proforma.rounded_total, Decimal('321512.00'))
        self.assertEqual(proforma.advance_amount, Decimal('86000.00'))
        self.assertEqual(proforma.balance_payable, Decimal('235512.00'))
        self.assertEqual(proforma.round_off_amount, Decimal('0.35'))

    def test_proforma_summary_order_and_zero_advance(self):
        """
        Verify calculation & total summary order:
        Amount -> Sub-Total -> Discount -> Sub-Total -> GST Add -> Total -> Rounding Off -> Advance (if entered) -> Balance Payable
        Even when advance is 0, 'BALANCE AMOUNT PAYABLE BY YOU' is present in PDF.
        """
        proforma = Proforma.objects.create(
            proforma_number='MCIPL/03/2627/0002/ZEROADV',
            proforma_date=date(2026, 3, 15),
            customer_name='ABC HYDRAULICS',
            customer_address='Mumbai, India',
            tax_rate=Decimal('12.00'),
            discount_percentage=Decimal('10.00'),
            advance_amount=Decimal('0.00'),
            round_off_enabled=True,
        )
        ProformaItem.objects.create(
            proforma=proforma,
            sr_no='01.',
            description='Test Hydraulic Valve',
            quantity=Decimal('1.00'),
            unit='No',
            unit_rate=Decimal('10000.00'),
        )
        proforma.recalculate_totals()

        self.assertEqual(proforma.subtotal, Decimal('10000.00'))
        self.assertEqual(proforma.discount_amount, Decimal('1000.00'))
        self.assertEqual(proforma.subtotal_after_discount, Decimal('9000.00'))
        self.assertEqual(proforma.tax_amount, Decimal('1080.00'))
        self.assertEqual(proforma.grand_total, Decimal('10080.00'))
        self.assertEqual(proforma.rounded_total, Decimal('10080.00'))
        self.assertEqual(proforma.balance_payable, Decimal('10080.00'))
        self.assertEqual(proforma.gst_display_label, 'GST ADD @12%')

        # Check detail preview
        detail_res = self.client.get(reverse('proforma_detail', args=[proforma.pk]))
        self.assertEqual(detail_res.status_code, 200)
        self.assertContains(detail_res, 'SUB-TOTAL')
        self.assertContains(detail_res, 'DISCOUNT 10% LESS')
        self.assertContains(detail_res, 'GST ADD @12%')
        self.assertContains(detail_res, 'BALANCE AMOUNT PAYABLE BY YOU')
        self.assertContains(detail_res, '10,080.00')

        # Check PDF preview
        pdf_res = self.client.get(reverse('proforma_pdf_preview', args=[proforma.pk]))
        self.assertEqual(pdf_res.status_code, 200)

    def test_quotation_item_description_optional(self):
        """Item description should NOT be compulsory when creating/editing a quotation."""
        create_url = reverse('quotation_create')
        post_data = {
            'company_name': 'MAXFLOW CONTROLS',
            'company_address': 'Plot No. 1, MIDC',
            'company_phone': '9876543210',
            'company_email': 'info@maxflowcontrols.com',
            'company_gstin': '27ABCDE1234F1Z5',
            'company_pan': 'ABCDE1234F',
            'quotation_number': 'QTN/OPTDESC/001',
            'quotation_date': '2026-10-07',
            'salutation': 'DEAR SIR,',
            'subject': 'QUOTATION FOR HYDRAULICS',
            'customer_name': 'TEST CUSTOMER WITHOUT DESC',
            'customer_address': 'Plot 10, MIDC Area, Pune',
            'price_terms': 'F.O.R. SITE',
            'freight_terms': 'EXTRA AT ACTUALS',
            'pf_terms': 'NIL',
            'discount_terms': 'NET',
            'payment_terms': '100% ADVANCE',
            'validity_terms': '30 DAYS',
            'signatory_company': 'FOR MAXFLOW CONTROLS',
            'tax_terms': 'AS INDICATED ABOVE (18% GST Extra)',
            'signatory_name': 'RAJ',
            'signatory_designation': 'MANAGER',
            'subtotal': '5000.00',
            'discount_percentage': '0.00',
            'pf_percentage': '0.00',
            'freight_amount': '0.00',
            'taxable_amount': '5000.00',
            'tax_amount': '900.00',
            'grand_total': '5900.00',
            # Formset management
            'items-TOTAL_FORMS': '2',
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            # Item 0: has quantity and rate, but NO description
            'items-0-id': '',
            'items-0-sr_no': '1',
            'items-0-description': '',
            'items-0-hsn_code': '84136090',
            'items-0-quantity': '2',
            'items-0-unit': 'NOS',
            'items-0-unit_rate': '2500.00',
            'items-0-gst_rate': '18',
            'items-0-delivery_schedule': 'EX STOCK',
            # Item 1: completely blank extra row where sr_no is set
            'items-1-id': '',
            'items-1-sr_no': '2',
            'items-1-description': '',
            'items-1-hsn_code': '',
            'items-1-quantity': '',
            'items-1-unit': 'NOS',
            'items-1-unit_rate': '',
            'items-1-gst_rate': '',
            'items-1-delivery_schedule': 'EX STOCK',
        }

        res = self.client.post(create_url, post_data)
        self.assertEqual(res.status_code, 302)

        q = Quotation.objects.get(quotation_number='QTN/OPTDESC/001')
        # Only the 1 substantive item should be created, extra blank row ignored
        self.assertEqual(q.items.count(), 1)
        item = q.items.first()
        self.assertEqual(item.description, '')
        self.assertEqual(item.quantity, Decimal('2.00'))
        self.assertEqual(item.unit_rate, Decimal('2500.00'))
        self.assertEqual(item.amount, Decimal('5000.00'))

        # Detail and PDF rendering must succeed with blank description
        detail_res = self.client.get(reverse('quotation_detail', args=[q.pk]))
        self.assertEqual(detail_res.status_code, 200)
        pdf_res = self.client.get(reverse('quotation_pdf_preview', args=[q.pk]))
        self.assertEqual(pdf_res.status_code, 200)

    def test_proforma_item_description_optional(self):
        """Item description should NOT be compulsory when editing a Proforma invoice."""
        proforma = Proforma.objects.create(
            proforma_number='PI/TEST/DESC/001',
            customer_name='CUSTOMER NO DESC',
            customer_address='Factory 2, MIDC',
            status='Draft',
        )
        edit_url = reverse('proforma_edit', args=[proforma.pk])
        post_data = {
            'proforma_number': 'PI/TEST/DESC/001',
            'customer_name': 'CUSTOMER NO DESC',
            'customer_address': 'Factory 2, MIDC',
            'status': 'Draft',
            'proforma_date': '2026-10-07',
            'tax_rate': '18.00',
            'items-TOTAL_FORMS': '2',
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            # Item 0 has no description
            'items-0-id': '',
            'items-0-sr_no': '01.',
            'items-0-description': '',
            'items-0-quantity': '3',
            'items-0-unit': 'No',
            'items-0-unit_rate': '1000.00',
            # Item 1 is an untouched extra row
            'items-1-id': '',
            'items-1-sr_no': '02.',
            'items-1-description': '',
            'items-1-quantity': '',
            'items-1-unit': 'No',
            'items-1-unit_rate': '',
        }
        res = self.client.post(edit_url, post_data)
        self.assertEqual(res.status_code, 302)

        proforma.refresh_from_db()
        self.assertEqual(proforma.items.count(), 1)
        item = proforma.items.first()
        self.assertEqual(item.description, '')
        self.assertEqual(item.quantity, Decimal('3.00'))
        self.assertEqual(item.unit_rate, Decimal('1000.00'))

    def test_quotation_date_auto_filled_with_current_date(self):
        """When creating a quotation, the date field must be automatically pre-populated with today's date."""
        from django.utils import timezone
        today_str = timezone.now().date().strftime('%Y-%m-%d')
        res = self.client.get(reverse('quotation_create'))
        self.assertEqual(res.status_code, 200)
        # Form initial must have today's date
        form = res.context['form']
        self.assertEqual(str(form.initial.get('quotation_date')), today_str)
        # Rendered HTML must contain value="YYYY-MM-DD"
        self.assertContains(res, f'value="{today_str}"')









