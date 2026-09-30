import io
import os
import re
import openpyxl
from decimal import Decimal
from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction
from quotations.models import Product


def read_file_safely(file_path):
    """
    Reads a file into bytes safely on Windows, even if it is currently open/locked
    by another application like Microsoft Excel (sharing violation / Errno 13).
    """
    try:
        with open(file_path, 'rb') as f:
            return f.read()
    except (PermissionError, OSError):
        pass

    if os.name == 'nt':
        try:
            import ctypes
            from ctypes import wintypes

            GENERIC_READ = 0x80000000
            FILE_SHARE_READ = 0x00000001
            FILE_SHARE_WRITE = 0x00000002
            FILE_SHARE_DELETE = 0x00000004
            OPEN_EXISTING = 3
            FILE_ATTRIBUTE_NORMAL = 0x80

            kernel32 = ctypes.WinDLL('kernel32', use_last_error=True)
            CreateFileW = kernel32.CreateFileW
            CreateFileW.argtypes = [
                wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE
            ]
            CreateFileW.restype = wintypes.HANDLE

            ReadFile = kernel32.ReadFile
            ReadFile.argtypes = [
                wintypes.HANDLE, wintypes.LPVOID, wintypes.DWORD,
                ctypes.POINTER(wintypes.DWORD), wintypes.LPVOID
            ]
            ReadFile.restype = wintypes.BOOL

            CloseHandle = kernel32.CloseHandle
            CloseHandle.argtypes = [wintypes.HANDLE]
            CloseHandle.restype = wintypes.BOOL

            GetFileSizeEx = kernel32.GetFileSizeEx
            GetFileSizeEx.argtypes = [wintypes.HANDLE, ctypes.POINTER(ctypes.c_int64)]
            GetFileSizeEx.restype = wintypes.BOOL

            handle = CreateFileW(
                str(file_path),
                GENERIC_READ,
                FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE,
                None,
                OPEN_EXISTING,
                FILE_ATTRIBUTE_NORMAL,
                None
            )

            if handle and handle != wintypes.HANDLE(-1).value:
                try:
                    size = ctypes.c_int64()
                    if GetFileSizeEx(handle, ctypes.byref(size)) and size.value > 0:
                        buf = ctypes.create_string_buffer(size.value)
                        read_bytes = wintypes.DWORD()
                        if ReadFile(handle, buf, size.value, ctypes.byref(read_bytes), None):
                            return buf.raw[:read_bytes.value]
                finally:
                    CloseHandle(handle)
        except Exception:
            pass

    with open(file_path, 'rb') as f:
        return f.read()


def load_workbook_safely(file_path):
    """Loads an openpyxl workbook safely from bytes, bypassing any Excel file locks."""
    file_bytes = read_file_safely(file_path)
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)


class Command(BaseCommand):
    help = 'Clears previous product items and imports all products from Excel (default: Product List.xlsx)'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='Product List.xlsx',
            help='Relative or absolute path to the product Excel file'
        )
        parser.add_argument(
            '--no-clear',
            action='store_true',
            help='Do not delete existing products before importing'
        )

    def handle(self, *args, **options):
        file_arg = options['file']
        if os.path.isabs(file_arg):
            excel_path = file_arg
        else:
            excel_path = os.path.join(settings.BASE_DIR, file_arg)

        if not os.path.exists(excel_path):
            self.stderr.write(self.style.ERROR(f"Excel file not found at: {excel_path}"))
            return

        self.stdout.write(f"Reading products safely from: {excel_path}")
        wb = load_workbook_safely(excel_path)
        sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            self.stderr.write(self.style.ERROR("Excel sheet is empty."))
            return

        # Find header row
        header_row_idx = None
        stock_idx = 0
        brand_idx = 1
        desc_idx = 2

        for idx, row in enumerate(rows[:10]):
            row_str = [str(c).strip().lower() if c is not None else '' for c in row]
            for col_idx, cell_text in enumerate(row_str):
                if 'stock item' in cell_text or 'item name' in cell_text or 'product' in cell_text:
                    header_row_idx = idx
                    stock_idx = col_idx
                    break
            if header_row_idx is not None:
                # also find brand and description indices if present
                for col_idx, cell_text in enumerate(row_str):
                    if 'brand' in cell_text or 'make' in cell_text or 'category' in cell_text:
                        brand_idx = col_idx
                    elif 'desc' in cell_text:
                        desc_idx = col_idx
                break

        if header_row_idx is None:
            # Fallback: assume row 2 (0-indexed 1) or row 1 (0-indexed 0)
            header_row_idx = 1 if len(rows) > 1 else 0

        data_rows = rows[header_row_idx + 1:]
        self.stdout.write(f"Detected header at row {header_row_idx + 1}. Data rows to process: {len(data_rows)}")

        products_to_create = []
        for r_num, row in enumerate(data_rows, start=header_row_idx + 2):
            if not row or not any(row):
                continue

            raw_name = row[stock_idx] if stock_idx < len(row) else None
            if raw_name is None:
                continue

            item_name = str(raw_name).strip()
            if not item_name or item_name.lower() in ['none', 'null', 'nan']:
                continue

            raw_brand = row[brand_idx] if brand_idx < len(row) else None
            raw_desc = row[desc_idx] if desc_idx < len(row) else None

            brand = str(raw_brand).strip() if raw_brand is not None else ''
            if not brand or brand.lower() in ['none', 'null', 'nan']:
                brand = 'GENERAL'

            desc = str(raw_desc).strip() if raw_desc is not None else ''
            if desc.lower() in ['none', 'null', 'nan']:
                desc = ''

            products_to_create.append(Product(
                category=brand,
                model_code=item_name,
                clean_code=re.sub(r'[^A-Za-z0-9]', '', item_name).upper(),
                description=desc,
                unit_rate=Decimal('0.00'),
                hsn_code='84136090',
                unit='NOS',
                gst_rate=Decimal('18.00'),
                is_active=True
            ))

        if not options.get('no_clear', False):
            old_count = Product.objects.count()
            self.stdout.write(f"Clearing {old_count} existing products as requested...")
            Product.objects.all().delete()

        with transaction.atomic():
            Product.objects.bulk_create(products_to_create, batch_size=1000)

        imported_count = len(products_to_create)
        self.stdout.write(
            self.style.SUCCESS(
                f"Successfully imported {imported_count} products from '{os.path.basename(excel_path)}' into catalog!"
            )
        )
